from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from stock_data import aggregate, resolve_stock, save_excel, validate_dates


class StockDataTests(unittest.TestCase):
    def setUp(self):
        self.data = pd.DataFrame({
            'Open': [10, 12, 20, 24], 'High': [15, 18, 26, 30],
            'Low': [8, 11, 19, 22], 'Close': [12, 17, 24, 29],
            'Volume': [100, 200, 300, 400],
        }, index=pd.to_datetime(['2025-01-30', '2025-01-31', '2025-02-03', '2025-02-04']))

    def test_weekly_and_monthly(self):
        for interval in ['주봉', '월봉']:
            rows = aggregate(self.data, date(2025, 1, 30), date(2025, 2, 4), interval)
            self.assertEqual(rows, [
                [date(2025, 1, 30), date(2025, 1, 31), 10, 18, 8, 17, 300],
                [date(2025, 2, 3), date(2025, 2, 4), 20, 30, 19, 29, 700]])

    def test_partial_period_inclusive_end(self):
        rows = aggregate(self.data, date(2025, 1, 31), date(2025, 2, 3), '주봉')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][2:], [12, 18, 11, 17, 200])
        self.assertEqual(rows[1][1], date(2025, 2, 3))

    def test_week_crossing_month(self):
        frame = self.data.iloc[:2].copy()
        frame.index = pd.to_datetime(['2025-01-31', '2025-02-01'])
        self.assertEqual(len(aggregate(frame, '2025-01-01', '2025-02-28', '주봉')), 1)
        self.assertEqual(len(aggregate(frame, '2025-01-01', '2025-02-28', '월봉')), 2)

    def test_invalid_dates_and_empty_range(self):
        for start, end in [('2025-02-30', '2025-03-01'), ('2025-02-01', '2025-01-01'), ('2099-01-01', '2099-01-02')]:
            with self.assertRaises(ValueError):
                validate_dates(start, end)
        with self.assertRaises(ValueError):
            aggregate(self.data, '2025-02-08', '2025-02-09', '일봉')

    def test_missing_and_duplicate_data(self):
        broken = self.data.copy()
        broken.iloc[0, 0] = float('nan')
        with self.assertRaises(ValueError):
            aggregate(broken, '2025-01-01', '2025-02-28', '일봉')
        with self.assertRaises(ValueError):
            aggregate(pd.concat([self.data, self.data]), '2025-01-01', '2025-02-28', '일봉')

    def test_name_resolution(self):
        listing = pd.DataFrame({'Name': ['삼성전자', '삼성전자우'], 'Code': ['005930', '005935']})
        self.assertEqual(resolve_stock('삼성전자', listing), ('005930', '삼성전자'))
        self.assertEqual(resolve_stock('005930', listing)[0], '005930')
        with self.assertRaisesRegex(ValueError, '검색 후보'):
            resolve_stock('삼성', listing)

    def test_excel_types_and_atomic_failure(self):
        rows = aggregate(self.data, '2025-01-01', '2025-02-28', '일봉')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'test.xlsx'
            args = (path, rows, '005930', '=literal', date(2025, 1, 1), date(2025, 2, 28), '일봉')
            save_excel(*args)
            with path.open('rb') as handle:
                wb = load_workbook(handle)
                self.assertEqual(wb.active.max_row, 5)
                self.assertEqual(wb.active['C2'].value, 10)
                self.assertTrue(wb.active['A2'].is_date)
                self.assertEqual(wb['조회 정보']['B3'].value, '005930')
                self.assertEqual(wb['조회 정보']['B2'].data_type, 's')
                self.assertEqual(wb.active.auto_filter.ref, 'A1:G5')
                wb.close()
            original = path.read_bytes()
            with patch('stock_data.os.replace', side_effect=PermissionError):
                with self.assertRaises(PermissionError):
                    save_excel(*args)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(len(list(Path(folder).iterdir())), 1)


if __name__ == '__main__':
    unittest.main()
