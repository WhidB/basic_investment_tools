"""Stock retrieval, calendar aggregation, and atomic Excel export."""
from datetime import date, datetime
from pathlib import Path
import os
import re
import tempfile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = ['Open', 'High', 'Low', 'Close', 'Volume']
LABELS = ['시가', '고가', '저가', '종가', '거래량']


def validate_dates(start, end):
    try:
        start, end = date.fromisoformat(start), date.fromisoformat(end)
    except ValueError as exc:
        raise ValueError('날짜를 YYYY-MM-DD 형식으로 입력해 주세요.') from exc
    if start > end:
        raise ValueError('시작일은 종료일보다 늦을 수 없습니다.')
    if end > date.today():
        raise ValueError('종료일은 오늘 또는 이전 날짜로 선택해 주세요.')
    return start, end


def resolve_stock(query, listing=None):
    query = query.strip()
    if not query:
        raise ValueError('주식 이름 또는 6자리 종목코드를 입력해 주세요.')
    if re.fullmatch(r'\d{4}[0-9A-HJ-NP-TV-Z][0-9KLMN]', query.upper()):
        return query.upper(), query.upper()
    if listing is None:
        import FinanceDataReader as fdr
        listing = fdr.StockListing('KRX')
    normalized = listing['Name'].astype(str).str.replace(' ', '').str.casefold()
    key = query.replace(' ', '').casefold()
    matches = listing[normalized == key]
    if len(matches) == 1:
        row = matches.iloc[0]
        return str(row['Code']).zfill(6), str(row['Name'])
    candidates = listing[normalized.str.contains(key, regex=False)]
    hints = ', '.join(f"{r['Name']} ({r['Code']})" for _, r in candidates.head(10).iterrows())
    raise ValueError('정확한 종목명이나 종목코드를 입력해 주세요.' + (f'\n검색 후보: {hints}' if hints else '\n해당 국내 상장 종목을 찾지 못했습니다.'))


def aggregate(data, start, end, interval):
    if interval not in ('일봉', '주봉', '월봉'):
        raise ValueError('일봉, 주봉, 월봉 중 선택해 주세요.')
    if data is None or data.empty:
        raise ValueError('조회 기간에 데이터가 없습니다. 종목코드와 거래일을 확인해 주세요.')
    missing = set(COLUMNS) - set(data.columns)
    if missing:
        raise ValueError(f'데이터 제공처 응답에 필수 항목이 없습니다: {missing}')
    data = data[COLUMNS].copy()
    data.index = pd.to_datetime(data.index).tz_localize(None).normalize()
    data = data.sort_index().loc[str(start):str(end)]
    if data.index.has_duplicates:
        raise ValueError('데이터에 중복 날짜가 있습니다. 다시 조회해 주세요.')
    data = data.apply(pd.to_numeric, errors='coerce')
    if data.empty:
        raise ValueError('조회 기간에 거래 데이터가 없습니다.')
    if data.isna().any().any():
        raise ValueError('제공 데이터에 누락된 가격 또는 거래량이 있어 저장하지 않았습니다.')
    # Some sources use zero OHLC during trading suspensions. Exclude such rows.
    data = data[(data[['Open', 'High', 'Low', 'Close']] > 0).all(axis=1)]
    if data.empty:
        raise ValueError('조회 기간에 유효한 거래 데이터가 없습니다.')
    groups = ([(None, data.loc[[idx]]) for idx in data.index] if interval == '일봉'
              else data.groupby(data.index.to_period('W-SUN' if interval == '주봉' else 'M')))
    rows = []
    for _, group in groups:
        rows.append([group.index[0].date(), group.index[-1].date(),
                     float(group.Open.iloc[0]), float(group.High.max()),
                     float(group.Low.min()), float(group.Close.iloc[-1]), int(group.Volume.sum())])
    return rows


def retrieve(query, start, end, interval):
    import FinanceDataReader as fdr
    code, name = resolve_stock(query)
    data = fdr.DataReader('NAVER:' + code, start, end)
    return code, name, aggregate(data, start, end, interval)


def save_excel(path, rows, code, name, start, end, interval):
    """Write temporary file beside destination, then replace only after success."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = '주가 데이터'
    ws.append(['기간 시작일', '기간 종료일', *LABELS])
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    for cells in ws.iter_rows(min_row=2):
        for cell in cells[:2]:
            cell.number_format = 'yyyy-mm-dd'
        for cell in cells[2:6]:
            cell.number_format = '#,##0.00'
        cells[6].number_format = '#,##0'
    info = wb.create_sheet('조회 정보')
    for row in [
        ['항목', '내용'], ['종목명', name], ['종목코드', code], ['시장 / 통화', '한국 / KRW'],
        ['조회 시작일', start], ['조회 종료일 (포함)', end], ['주기', interval],
        ['저장 시각 (PC 현지 시간)', datetime.now().isoformat(timespec='seconds')],
        ['출처', 'NAVER Finance / FinanceDataReader'],
        ['출처 URL', f'https://finance.naver.com/item/sise_day.naver?code={code}'],
        ['집계 기준', '선택 기간 안의 거래일만 집계. 주봉: 월~일, 월봉: 달력 월. 경계의 주/월은 부분 봉.'],
        ['가격 기준', '제공처의 일별 OHLCV 사용. 별도 가격 보정 없음. 과거 가격은 제공처 정책에 따라 조정될 수 있음.'],
        ['제외 기준', '시가/고가/저가/종가가 0 이하인 행 제외 (거래정지 등).'],
        ['기간 표시', '각 봉에 실제 포함된 첫 거래일과 마지막 거래일. 당일 데이터는 장중 변동 가능.'],
    ]:
        info.append(row)
    # All metadata is literal text, even if a supplied name starts with '='.
    for row in info:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for sheet in wb:
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='16324F')
            cell.font = Font(color='FFFFFF', bold=True)
        sheet.row_dimensions[1].height = 26
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 19
    info.column_dimensions['A'].width = 30
    info.column_dimensions['B'].width = 100
    for row in info.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
        info.row_dimensions[row[0].row].height = 32
    fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(fd)
    try:
        wb.save(temp)
        os.replace(temp, path)
    finally:
        wb.close()
        if os.path.exists(temp):
            os.unlink(temp)
